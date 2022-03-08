# vgmc_pp_maker.py
from pathlib import Path
ROOT = Path(__file__).parent.absolute()

from pptx import Presentation
from openpyxl import load_workbook

from argparse import ArgumentParser
from lib.custom_exceptions import *
from lib.misc import determine_parameters, save_presentation
from lib.slides import *


# Parse the arguments provided to argparse.
PARSER = ArgumentParser(description="Use this file to create a PowerPoint for the Florida Southern College Computer Science Club's Video Game Music Competiton (FSCCSCVGMC).")
PARSER.add_argument('-d', '--default', action="store_true", help='Create a default Competiton of 5 Rounds each with 10 Tracks')
PARSER.add_argument('-r', '--rounds', type=int, nargs='?', help='Provide a custom number of rounds')
PARSER.add_argument('-t', '--tracks', type=int, nargs='?', help='Provide a custom number of tracks per round')
PARSER.add_argument('-nc', '--nochallenge', action='store_true', help="This flag disables the final round being a Challege Round")
ARGS = PARSER.parse_args()

def main():
	
	# Determine how many Rounds and Tracks per Round there will be.
	num_rounds, num_tracks, challenge = determine_parameters(ARGS)
	print(f"\033[96mThere will be \033[93m{num_rounds}\033[96m Round{'s Each' if num_rounds > 1 else ''} with \033[93m{num_tracks}\033[96m Tracks.\nRound {num_rounds} will {'also' if challenge else 'not'} be a Challenge Round.\n\033[0m")

	# Create the PowerPoint using the slide template.
	prs = Presentation(ROOT / "templates" / "master_copy.pptx")

	# Read in the answer key for later
	wb = load_workbook(ROOT / 'tracks' / 'song_info.xlsx')
	answer_sheet = wb.active

	# Create the Title and Rules Slides
	rules_slide(prs, num_rounds, num_tracks)

	# Create the slides for each round.
	for round_num in range(1, num_rounds + 1):

		print(f"\033[96mGenerating Round {round_num}...\033[0m")


		challenge_round = (challenge and round_num == num_rounds)

		round_slide(prs, round_num, challenge_round)

		if challenge_round:
			rules_slide(prs, num_rounds, num_tracks, True)

		for track_num in range(1, num_tracks + 1):
			track_slide(prs, round_num, track_num)

		review_slide(prs, round_num, num_tracks)

		answer_slide(prs, answer_sheet, round_num, num_tracks, challenge_round)


	# Save the Presentation
	save_presentation(prs)


if __name__ == '__main__':
	print('\n\033[92mRunning vgmc_pp_maker.py\033[0m')

	try:
		main()

	except Exception as e:
		print(f'\033[91m{e}\033[0m')